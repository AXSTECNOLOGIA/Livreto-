"""Teste ponta a ponta do pipeline com um FIXTURE SINTÉTICO.

Os arquivos gerados aqui imitam apenas o layout (nomes de arquivo, colunas,
codificação latin-1, delimitador ';') dos arquivos do TSE. Os números são
inventados para exercitar cada regra de classificação e NÃO representam
nenhum resultado eleitoral. Nada deste teste é gravado fora de um diretório
temporário.

Execução:  python -m unittest discover -s tests -v
"""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

RAIZ = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("auditoria", RAIZ / "Auditoria_TSE_2026.py")
aud = importlib.util.module_from_spec(spec)
sys.modules["auditoria"] = aud
spec.loader.exec_module(aud)

COLS_DET = ["DT_GERACAO", "HH_GERACAO", "ANO_ELEICAO", "CD_TIPO_ELEICAO", "NM_TIPO_ELEICAO", "NR_TURNO",
            "CD_ELEICAO", "DS_ELEICAO", "DT_ELEICAO", "TP_ABRANGENCIA", "SG_UF", "SG_UE", "NM_UE",
            "CD_MUNICIPIO", "NM_MUNICIPIO", "NR_ZONA", "NR_SECAO", "CD_CARGO", "DS_CARGO", "QT_APTOS",
            "QT_COMPARECIMENTO", "QT_ABSTENCOES", "QT_VOTOS_NOMINAIS", "QT_VOTOS_BRANCOS", "QT_VOTOS_NULOS",
            "QT_VOTOS_LEGENDA", "QT_VOTOS_ANULADOS_APU_SEP"]
COLS_VS = ["DT_GERACAO", "ANO_ELEICAO", "NR_TURNO", "SG_UF", "CD_MUNICIPIO", "NM_MUNICIPIO", "NR_ZONA",
           "NR_SECAO", "CD_CARGO", "DS_CARGO", "NR_VOTAVEL", "NM_VOTAVEL", "QT_VOTOS"]
COLS_BU = ["DT_GERACAO", "NR_TURNO", "SG_UF", "CD_MUNICIPIO", "NM_MUNICIPIO", "NR_ZONA", "NR_SECAO",
           "NR_LOCAL_VOTACAO", "CD_CARGO_PERGUNTA", "DS_CARGO_PERGUNTA", "QT_APTOS", "QT_COMPARECIMENTO",
           "QT_ABSTENCOES", "DS_TIPO_URNA", "DS_TIPO_VOTAVEL", "NR_VOTAVEL", "NM_VOTAVEL", "QT_VOTOS",
           "NR_URNA_EFETIVADA", "DS_AGREGADAS"]
COLS_ELV = ["AA_ELEICAO", "NR_TURNO", "SG_UF", "CD_MUNICIPIO", "NM_MUNICIPIO", "NR_ZONA", "NR_SECAO",
            "CD_TIPO_SECAO_AGREGADA", "DS_TIPO_SECAO_AGREGADA", "NR_SECAO_PRINCIPAL", "NR_LOCAL_VOTACAO",
            "NM_LOCAL_VOTACAO", "DS_TIPO_LOCAL", "QT_ELEITOR_SECAO"]

MUNICIPIOS = [("AC", 1120, "RIO BRANCO", 40), ("AC", 1392, "CRUZEIRO DO SUL", 40), ("AC", 1007, "BUJARI", 5),
              ("DF", 97012, "BRASÍLIA", 60)]
CANDIDATOS = {1: [13, 22, 12, 15], 2: [13, 22]}


def csv_bytes(cols, linhas) -> bytes:
    buf = io.StringIO()
    wr = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\n")
    wr.writerow(cols)
    wr.writerows(linhas)
    return buf.getvalue().encode("latin-1")


def gravar_zip(caminho: Path, membros: dict[str, bytes]) -> None:
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as zf:
        for nome, dados in membros.items():
            zf.writestr(nome, dados)


def gerar_fixture(destino: Path) -> dict:
    rng = np.random.default_rng(42)
    secoes = []  # (uf, mun, nome, zona, secao, aptos, comp1, comp2)
    for uf, mun, nome, n in MUNICIPIOS:
        for i in range(n):
            aptos = int(rng.integers(300, 400))
            c1 = int(round(aptos * 0.8 + rng.normal(0, 4)))
            c2 = int(round(c1 + rng.normal(0, 4)))
            secoes.append(dict(uf=uf, mun=mun, nome=nome, zona=1, secao=101 + i, a1=aptos, a2=aptos,
                               c1=c1, c2=c2, ag1="", ag2="", local=1000 + i // 5))
    idx = {(s["uf"], s["mun"], s["secao"]): s for s in secoes}
    S = lambda n: idx[("AC", 1120, n)]  # noqa: E731

    S(101)["c1"] = int(S(101)["a1"] * 0.65)
    S(101)["c2"] = int(S(101)["c1"] * 1.30)                       # aumento atípico, BU confere -> C2
    S(102)["c2"] = S(102)["a2"] + 7                               # comparecimento > aptos (no BU também) -> C4
    S(103)["bu_c1_delta"] = 5                                     # BU difere do detalhe no 1T -> C4
    S(104)["so_1t"] = True                                        # agregada à 105 no 2T -> C3
    S(105)["a2"] = S(105)["a1"] + S(104)["a1"]
    S(105)["c2"] = S(105)["c1"] + S(104)["c1"]
    S(105)["ag2"] = "104"                                         # alteração de agregação -> C3
    S(106)["c1"] = 0                                              # comparecimento 1T zero -> C5
    S(107)["a2"] = S(107)["a1"] + 50                              # variação de aptos sem explicação -> C5
    S(107)["c2"] = S(107)["c1"] + 40
    S(108)["c1"] = int(S(108)["a1"] * 0.65)
    S(108)["c2"] = int(S(108)["c1"] * 1.35)                       # atípica, sem BU no 2T -> C5
    S(108)["sem_bu_2t"] = True

    det = {"AC": [], "DF": []}
    vs = {("AC"): [], "DF": []}
    bu = {(u, t): [] for u in ("AC", "DF") for t in (1, 2)}
    elv = []
    for s in secoes:
        for t in (1, 2):
            if t == 2 and s.get("so_1t"):
                continue
            a, c = s[f"a{t}"], s[f"c{t}"]
            brancos, nulos = c // 40, c // 30
            validos = c - brancos - nulos
            cand = CANDIDATOS[t]
            partes = [validos // len(cand)] * len(cand)
            partes[0] += validos - sum(partes)
            base = ["26/10/2026", "10:00:00", "2026", "2", "Eleição Ordinária", str(t), "9999",
                    "Eleição Geral Federal 2026", "", "F", s["uf"], s["uf"], s["uf"], str(s["mun"]), s["nome"],
                    str(s["zona"]), str(s["secao"])]
            det[s["uf"]].append(base + ["1", "Presidente", a, c, a - c, validos, brancos, nulos, 0, 0])
            if t == 1:  # outro cargo com números diferentes: precisa ser ignorado
                det[s["uf"]].append(base + ["3", "Governador", a, c - 7, a - c + 7, validos - 7, brancos, nulos, 0, 0])
            for nr, q in list(zip(cand, partes)) + [(95, brancos), (96, nulos)]:
                vs[s["uf"]].append(["26/10/2026", "2026", t, s["uf"], s["mun"], s["nome"], s["zona"], s["secao"],
                                    1, "Presidente", nr, f"VOTÁVEL {nr}", q])
            if t == 2 and s.get("sem_bu_2t"):
                continue
            cb = c + (s.get("bu_c1_delta", 0) if t == 1 else 0)
            tipos = [("Nominal", nr, q) for nr, q in zip(cand, partes)] + [("Branco", 95, brancos), ("Nulo", 96, nulos)]
            for ds, nr, q in tipos:
                bu[(s["uf"], t)].append(["26/10/2026", t, s["uf"], s["mun"], s["nome"], s["zona"], s["secao"],
                                         s["local"], 1, "Presidente", a, cb, a - cb, "Apurada", ds, nr,
                                         f"VOTÁVEL {nr}", q, 2000000 + s["secao"], s[f"ag{t}"] or "#NULO#"])
        for t in (1, 2):
            agregada = t == 2 and s.get("so_1t")
            elv.append([2026, t, s["uf"], s["mun"], s["nome"], s["zona"], s["secao"], 2 if agregada else 1,
                        "AGREGADA" if agregada else "PRINCIPAL", 105 if agregada else s["secao"], s["local"],
                        "ESCOLA ESTADUAL", "CONVENCIONAL", s["a1"]])

    gravar_zip(destino / "detalhe_votacao_secao_2026.zip", {
        "detalhe_votacao_secao_2026_AC.csv": csv_bytes(COLS_DET, det["AC"]),
        "detalhe_votacao_secao_2026_DF.csv": csv_bytes(COLS_DET, det["DF"]),
        # consolidado nacional: deve ser ignorado para não duplicar seções
        "detalhe_votacao_secao_2026_BRASIL.csv": csv_bytes(COLS_DET, det["AC"] + det["DF"]),
        "leiame.pdf": b"%PDF-1.4",
    })
    for u in ("AC", "DF"):
        gravar_zip(destino / f"votacao_secao_2026_{u}.zip", {f"votacao_secao_2026_{u}.csv": csv_bytes(COLS_VS, vs[u])})
        for t in (1, 2):
            gravar_zip(destino / f"bweb_{t}t_{u}_2710202612{t}0.zip",
                       {f"bweb_{t}t_{u}.csv": csv_bytes(COLS_BU, bu[(u, t)])})
    # versão antiga do BU (deve ser preterida pela mais recente)
    gravar_zip(destino / "bweb_1t_AC_011020261200.zip",
               {"bweb_1t_AC.csv": b"lixo;sem;esquema\n1;2;3\n"})
    gravar_zip(destino / "eleitorado_local_votacao_2026.zip",
               {"eleitorado_local_votacao_2026.csv": csv_bytes(COLS_ELV, elv)})
    return idx


class TestePipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        base = Path(cls.tmp.name)
        cls.entrada, cls.saida = base / "entrada", base / "saida"
        cls.entrada.mkdir()
        gerar_fixture(cls.entrada)
        cls.rc = aud.main(["--entrada", str(cls.entrada), "--saida", str(cls.saida), "--ufs", "AC,DF"])
        cls.w = pd.read_csv(cls.saida / "secoes_consolidadas.csv.gz", sep=";")
        cls.manifesto = json.loads((cls.saida / "manifesto_fontes.json").read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def secao(self, uf, mun, n):
        r = self.w[(self.w.SG_UF == uf) & (self.w.CD_MUNICIPIO == mun) & (self.w.NR_SECAO == n)]
        self.assertEqual(len(r), 1, f"seção {uf}-{mun}-{n} deveria aparecer uma vez")
        return r.iloc[0]

    def test_escopo_parcial_quando_ufs_restritas(self):
        self.assertEqual(self.rc, 1)
        self.assertEqual(self.manifesto["escopo"]["status"], "PARCIAL")

    def test_nao_duplica_brasil_nem_soma_outro_cargo(self):
        self.assertFalse(self.w.duplicated(aud.KEY).any())
        self.assertEqual(int(self.w.presente_1t.sum()), sum(m[3] for m in MUNICIPIOS))
        s = self.secao("AC", 1120, 110)
        self.assertEqual(s.comparecimento_1t + s.abstencoes_1t, s.aptos_1t)

    def test_calculos(self):
        s = self.secao("AC", 1120, 101)
        self.assertEqual(s.delta_comparecimento, s.comparecimento_2t - s.comparecimento_1t)
        self.assertAlmostEqual(s.variacao_pct, 100 * s.delta_comparecimento / s.comparecimento_1t)
        self.assertAlmostEqual(s.delta_taxa_pp, 100 * (s.comparecimento_2t / s.aptos_2t - s.comparecimento_1t / s.aptos_1t))

    def test_classificacoes_injetadas(self):
        esperado = {101: "C2", 102: "C4", 103: "C4", 104: "C3", 105: "C3", 106: "C5", 107: "C5", 108: "C5"}
        obtido = {n: self.secao("AC", 1120, n).classificacao for n in esperado}
        self.assertEqual(obtido, esperado)
        self.assertEqual(self.secao("AC", 1120, 104).situacao_correspondencia, "SOMENTE_1T")
        self.assertEqual(self.secao("AC", 1120, 105).situacao_correspondencia, "ALTERACAO_AGREGACAO")
        self.assertEqual(self.secao("AC", 1120, 106).situacao_correspondencia, "COMPARECIMENTO_1T_ZERO")
        self.assertTrue(self.secao("AC", 1120, 108).atipica)

    def test_secoes_normais_sao_c1_com_bu(self):
        df = self.w[self.w.SG_UF == "DF"]
        self.assertTrue((df.classificacao == "C1").all(), df.classificacao.value_counts().to_dict())
        self.assertTrue(df.bu_conferido.all())

    def test_grupo_alternativo_para_municipio_pequeno(self):
        s = self.secao("AC", 1007, 101)
        self.assertTrue(str(s.grupo_referencia).startswith("UF_FAIXA:AC"))

    def test_versao_mais_recente_do_bu(self):
        bu_ac1 = [f for f in self.manifesto["fontes"] if f["tipo"] == "bweb" and f["uf"] == "AC" and f["turno"] == 1]
        self.assertEqual(len(bu_ac1), 1)
        self.assertEqual(bu_ac1[0]["status"], "PROCESSADO")
        self.assertIn("2710202612", bu_ac1[0]["caminho"])

    def test_manifesto_sha256(self):
        for f in self.manifesto["fontes"]:
            if f["caminho"]:
                self.assertEqual(f["sha256"], hashlib.sha256(Path(f["caminho"]).read_bytes()).hexdigest())

    def test_entregaveis(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.saida / "Auditoria_Seções_2026.xlsx", read_only=True)
        self.assertEqual(wb.sheetnames, ["Resumo Nacional", "Resumo por UF", "Todas as Seções Comparáveis",
                                         "Aumentos Atípicos", "Inconsistências Matemáticas",
                                         "Validações dos Boletins", "Seções Não Comparáveis", "Fontes e Metodologia"])
        self.assertGreater((self.saida / "Auditoria_TSE_2026.pdf").stat().st_size, 2000)


class TesteSemDados(unittest.TestCase):
    def test_sem_arquivos_nao_inventa_resultado(self):
        with tempfile.TemporaryDirectory() as d:
            vazio, saida = Path(d) / "vazio", Path(d) / "saida"
            vazio.mkdir()
            rc = aud.main(["--entrada", str(vazio), "--saida", str(saida)])
            self.assertEqual(rc, 2)
            m = json.loads((saida / "manifesto_fontes.json").read_text(encoding="utf-8"))
            self.assertEqual(m["escopo"]["status"], "NÃO EXECUTADA")
            self.assertFalse((saida / "secoes_consolidadas.csv.gz").exists())


if __name__ == "__main__":
    unittest.main()
